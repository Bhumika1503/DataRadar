import io
import pandas as pd
import pymupdf

from docx import Document


SUPPORTED_TYPES = [
    "csv",
    "xlsx",
    "pdf",
    "docx",
    "txt",
]


def extract_file(uploaded_file):
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    result = {
        "file_type": filename.rsplit(".", 1)[-1],
        "dataframe": None,
        "text": "",
        "tables": [],
    }

    if filename.endswith(".csv"):

        result["dataframe"] = pd.read_csv(
            io.BytesIO(file_bytes)
        )

    elif filename.endswith(".xlsx"):

        result["dataframe"] = pd.read_excel(
            io.BytesIO(file_bytes)
        )

    elif filename.endswith(".txt"):

        try:
            result["text"] = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            result["text"] = file_bytes.decode("utf-8-sig")

    elif filename.endswith(".pdf"):

        with pymupdf.open(
            stream=file_bytes,
            filetype="pdf"
        ) as document:

            pages = []
            extracted_tables = []

            for page in document:

                pages.append(
                    page.get_text()
                )

                try:
                    tables = page.find_tables().tables
                except Exception:
                    tables = []

                for table in tables:

                    rows = table.extract()

                    if not rows:
                        continue

                    headers = rows[0]
                    data = rows[1:]

                    headers = [
                        str(header).strip()
                        if header is not None
                        else f"column_{i + 1}"
                        for i, header in enumerate(headers)
                    ]

                    if len(set(headers)) != len(headers):

                        headers = [
                            f"column_{i + 1}"
                            for i in range(len(headers))
                        ]

                    if data:

                        extracted_tables.append(
                            pd.DataFrame(
                                data,
                                columns=headers
                            )
                        )

            result["text"] = "\n\n".join(pages)
            result["tables"] = extracted_tables

    elif filename.endswith(".docx"):

        document = Document(
            io.BytesIO(file_bytes)
        )

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        result["text"] = "\n".join(paragraphs)

        extracted_tables = []

        for table in document.tables:

            rows = [
                [
                    cell.text.strip()
                    for cell in row.cells
                ]
                for row in table.rows
            ]

            if not rows:
                continue

            headers = rows[0]
            data = rows[1:]

            headers = [
                header if header else f"column_{i + 1}"
                for i, header in enumerate(headers)
            ]

            if len(set(headers)) != len(headers):

                headers = [
                    f"column_{i + 1}"
                    for i in range(len(headers))
                ]

            extracted_tables.append(
                pd.DataFrame(
                    data,
                    columns=headers
                )
            )

        result["tables"] = extracted_tables

    else:

        raise ValueError(
            f"Unsupported file type: {filename}"
        )

    return result


def get_excel_sheets(uploaded_file):

    file_bytes = uploaded_file.getvalue()

    excel = pd.ExcelFile(
        io.BytesIO(file_bytes)
    )

    try:
        return excel.sheet_names
    finally:
        excel.close()